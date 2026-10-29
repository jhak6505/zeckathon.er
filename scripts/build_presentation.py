"""
FINHYPE: Upgraded Professional 9-Slide Presentation Generator (v2)
Zeckathon 2026 · Problem Statement 3: Mutual Fund Cash-Flow Forecasting & Liquidity Buffer Sizing
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette (Dark Navy Theme with Financial Accents) ───────────────────
BG_COLOR = RGBColor(11, 19, 43)        # #0B132B Deep Navy
SURFACE_COLOR = RGBColor(15, 23, 42)   # #0F172A Slate Dark Navy
CARD_BG = RGBColor(30, 41, 59)         # #1E293B Card Background
CARD_BG_ALT = RGBColor(24, 34, 53)     # #182235 Darker Card
CARD_BG_LIGHT = RGBColor(40, 53, 76)   # #28354C Lighter Card Header
CARD_HIGHLIGHT = RGBColor(23, 37, 84)  # #172554 Navy Highlight Card

BORDER_COLOR = RGBColor(51, 65, 85)    # #334155 Slate Border
BORDER_BLUE = RGBColor(59, 130, 246)   # #3B82F6 Blue Accent Border
BORDER_CYAN = RGBColor(14, 165, 233)   # #0EA5E9 Sky Blue Border
BORDER_AMBER = RGBColor(245, 158, 11)  # #F59E0B Amber Border
BORDER_GREEN = RGBColor(16, 185, 129)  # #10B981 Emerald Border

TEXT_PRIMARY = RGBColor(248, 250, 252)   # #F8FAFC Pure Crisp White
TEXT_SECONDARY = RGBColor(203, 213, 225) # #CBD5E1 High-Contrast Light Slate
TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8 Subtle Muted Slate

ACCENT_BLUE = RGBColor(96, 165, 250)   # #60A5FA Primary Blue
ACCENT_CYAN = RGBColor(56, 189, 248)   # #38BDF8 Cyan Metric Highlight
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399 Positive / Inflow / Winner
ACCENT_AMBER = RGBColor(251, 191, 36)  # #FBBF24 Warning / Moderate Risk
ACCENT_RED = RGBColor(248, 113, 113)   # #F87171 Negative / Downside Stress

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"


def create_base_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def set_slide_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.color.rgb = BG_COLOR


def add_slide_header(slide, title_text: str, category_tag: str):
    # Category Tracker Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.733), Inches(0.32))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_tag.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(10.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.733), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_PRIMARY

    # Accent divider line under header
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.30), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER_COLOR
    line.line.color.rgb = BORDER_COLOR


def add_slide_footer(slide, current_slide: int, total_slides: int = 9, source_note: str = None):
    # Footer divider
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.015))
    div.fill.solid()
    div.fill.fore_color.rgb = BORDER_COLOR
    div.line.color.rgb = BORDER_COLOR

    # Left: Project info
    left_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(5.8), Inches(0.35))
    tf_l = left_box.text_frame
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "FINHYPE | Zeckathon 2026 · Problem Statement 3 | Origin: Sep 2026 · Target: Oct–Dec 2026"
    p_l.font.name = FONT_BODY
    p_l.font.size = Pt(9)
    p_l.font.color.rgb = TEXT_MUTED

    # Center/Right: Source note if present
    if source_note:
        src_box = slide.shapes.add_textbox(Inches(6.6), Inches(6.92), Inches(5.3), Inches(0.35))
        tf_s = src_box.text_frame
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = f"Source: {source_note}"
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(8.5)
        p_s.font.italic = True
        p_s.font.color.rgb = TEXT_MUTED
        p_s.alignment = PP_ALIGN.RIGHT

    # Far Right: Slide Number
    num_box = slide.shapes.add_textbox(Inches(12.0), Inches(6.92), Inches(0.533), Inches(0.35))
    tf_n = num_box.text_frame
    tf_n.margin_left = tf_n.margin_right = tf_n.margin_top = tf_n.margin_bottom = 0
    p_n = tf_n.paragraphs[0]
    p_n.text = f"{current_slide} / {total_slides}"
    p_n.font.name = FONT_BODY
    p_n.font.size = Pt(9.5)
    p_n.font.bold = True
    p_n.font.color.rgb = ACCENT_BLUE
    p_n.alignment = PP_ALIGN.RIGHT


def add_card(slide, left: float, top: float, width: float, height: float,
             bg_color=CARD_BG, border_color=BORDER_COLOR):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card


def format_bullet_point(tf, title: str, body: str, is_first: bool = False,
                        title_color=TEXT_PRIMARY, body_color=TEXT_SECONDARY,
                        font_size_title=12, font_size_body=11, space_after=10):
    p = tf.paragraphs[0] if is_first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    
    if title:
        r_title = p.add_run()
        r_title.text = f"{title}: "
        r_title.font.name = FONT_HEADING
        r_title.font.bold = True
        r_title.font.size = Pt(font_size_title)
        r_title.font.color.rgb = title_color

    r_body = p.add_run()
    r_body.text = body
    r_body.font.name = FONT_BODY
    r_body.font.size = Pt(font_size_body)
    r_body.font.color.rgb = body_color


# ── Slide 1: Title and Value Proposition ─────────────────────────────────────
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Hackathon Banner Tag
    tag_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.95), Inches(11.333), Inches(0.35))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "ZECKATHON 2026 · PROBLEM STATEMENT 3: MUTUAL FUND CASH FLOW & LIQUIDITY SIZING"
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(11.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN

    # Main Project Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(11.333), Inches(1.1))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = "FINHYPE"
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_PRIMARY

    # Subtitle / Core Theme
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.50), Inches(11.333), Inches(0.55))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Quantitative Treasury Liquidity Forecasting & Risk-Calibrated Buffer Sizing"
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(22)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_BLUE

    # Value Proposition Callout Card
    add_card(slide, 1.0, 3.25, 11.333, 1.25, bg_color=CARD_BG, border_color=BORDER_BLUE)
    vp_box = slide.shapes.add_textbox(Inches(1.25), Inches(3.38), Inches(10.833), Inches(1.0))
    tf_vp = vp_box.text_frame
    tf_vp.word_wrap = True
    p_vp = tf_vp.paragraphs[0]
    p_vp.text = (
        "An end-to-end quantitative platform that pre-positions Indian Liquid and Overnight mutual fund liquidity "
        "against cyclical corporate tax drains through leak-free probabilistic forecasting and empirical downside deficit sizing."
    )
    p_vp.font.name = FONT_BODY
    p_vp.font.size = Pt(14)
    p_vp.font.color.rgb = TEXT_PRIMARY

    # 3 Strategic Pillar Cards
    pillars = [
        ("01 / Temporal Integrity", "90 continuous AMFI monthly statements (Apr 2019 – Sep 2026), zero-leakage shifted feature engineering, and 154 passing unit & integration tests.", ACCENT_CYAN),
        ("02 / Empirical Hierarchy", "Four forecasting models evaluated across 24 expanding rolling backtests; SARIMAX and Seasonal-Naive selected based on empirical horizon skill.", ACCENT_BLUE),
        ("03 / Solvency Buffer Engine", "Two-tier liquidity sizing combining operational floor with peak cumulative downside deficit under conservative comonotonic stress.", ACCENT_GREEN),
    ]

    col_width = 3.6
    gap = 0.266
    for idx, (p_title, p_desc, col_accent) in enumerate(pillars):
        left = 1.0 + idx * (col_width + gap)
        add_card(slide, left, 4.70, col_width, 1.95, bg_color=CARD_BG, border_color=BORDER_COLOR)
        
        # Pillar header
        h_box = slide.shapes.add_textbox(Inches(left + 0.2), Inches(4.85), Inches(col_width - 0.4), Inches(0.4))
        tf_h = h_box.text_frame
        p_h = tf_h.paragraphs[0]
        p_h.text = p_title
        p_h.font.name = FONT_HEADING
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = col_accent

        # Pillar desc
        d_box = slide.shapes.add_textbox(Inches(left + 0.2), Inches(5.30), Inches(col_width - 0.4), Inches(1.25))
        tf_d = d_box.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = p_desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_SECONDARY

    add_slide_footer(slide, 1, 9, source_note="AMFI Official Data; Forecast Origin: September 2026")


# ── Slide 2: The Liquidity Planning Problem (Corrected Turnover Figures) ─────
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "The Liquidity Sizing Problem in Indian Debt Mutual Funds", "Problem Context")

    col_w = 3.7
    gap = 0.316

    # Card 1: Advance-Tax Cycles
    c1_left = 0.8
    add_card(slide, c1_left, 1.50, col_w, 5.15)
    tb1 = slide.shapes.add_textbox(Inches(c1_left + 0.22), Inches(1.68), Inches(col_w - 0.44), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "Corporate Advance-Tax Cycles"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(12)

    format_bullet_point(tf1, "Predictable Cash Drains", "Indian corporate treasuries treat Liquid and Overnight fund units as operational cash accounts. Every quarter (Mar, Jun, Sep, Dec), corporations withdraw tens of thousands of crores to pay advance tax installments.", space_after=10)
    format_bullet_point(tf1, "Post-Quarter Inflow Surges", "Following advance-tax deadlines, corporate liquidity returns in the subsequent month (Apr, Jul, Oct, Jan) as operating revenues replenish treasury balances.", space_after=10)
    format_bullet_point(tf1, "December 2026 Projection", "SARIMAX forecasts a severe net drain of -₹50,618.62 Cr for Liquid Funds in December 2026 (lower 80% bound: -₹89,235.07 Cr).", space_after=6, body_color=ACCENT_RED)

    # Card 2: Velocity Disparity (Audited Turnover Numbers)
    c2_left = c1_left + col_w + gap
    add_card(slide, c2_left, 1.50, col_w, 5.15)
    tb2 = slide.shapes.add_textbox(Inches(c2_left + 0.22), Inches(1.68), Inches(col_w - 0.44), Inches(4.8))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "Liquid vs. Overnight Velocity"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(12)

    format_bullet_point(tf2, "Liquid Funds (~₹6.26L Cr AUM)", "Average monthly redemptions are ~₹4.64 lakh Cr (~104% of AUM; median 74%), turning over ~1.0x of fund assets monthly.", space_after=10)
    format_bullet_point(tf2, "Overnight Funds (~₹84K Cr AUM)", "Average monthly redemptions are ~₹4.22 lakh Cr (~545% of AUM; median 504%), turning over 5.45x of fund assets monthly.", space_after=10)
    format_bullet_point(tf2, ">5x Velocity Gap", "Overnight Funds turn over their asset base more than 5 times faster than Liquid Funds, making them acutely sensitive to gross repurchase volatility.", space_after=6, body_color=ACCENT_AMBER)

    # Card 3: The Asymmetric Cost of Error
    c3_left = c2_left + col_w + gap
    add_card(slide, c3_left, 1.50, col_w, 5.15)
    tb3 = slide.shapes.add_textbox(Inches(c3_left + 0.22), Inches(1.68), Inches(col_w - 0.44), Inches(4.8))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "The Asymmetric Cost of Error"
    p.font.name = FONT_HEADING
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_after = Pt(12)

    format_bullet_point(tf3, "Under-Buffering (Solvency Risk)", "Underestimating net outflows forces distress fire-sales of commercial papers/T-bills or costly emergency borrowing in repo markets at punitive rates, triggering reputation damage.", space_after=10)
    format_bullet_point(tf3, "Over-Buffering (Yield Drag)", "Holding excessive uninvested cash depresses fund yields. In basis-point competitive markets, a 10 bps yield drag triggers massive institutional flight to competing AMCs.", space_after=10)
    format_bullet_point(tf3, "The Core Technical Gap", "Point forecasting (E[Y]) tells expected flows, not required solvency cash. Sizing requires empirical downside error calibration across multi-month horizons.", space_after=6, body_color=ACCENT_CYAN)

    add_slide_footer(slide, 2, 9, source_note="AMFI Monthly Statements (Apr 2019 – Sep 2026); data/processed/fund_flows_monthly.csv")


# ── Slide 3: Data Sources and Validation (Corrected Repo Rate 4.0% - 6.5%) ───
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Data Foundation: 90 Continuous Months & Multi-Stage Validation", "Data Engineering")

    col_w = 2.75
    gap = 0.244
    pillars = [
        ("AMFI Monthly Statements", "90 consecutive monthly statements from Apr 2019 to Sep 2026 (180 category records) across Liquid and Overnight funds.", ACCENT_CYAN),
        ("Format Normalization", "Automated regex parser standardizes 7+ years of AMFI layout shifts, header renames, and Indian number formatting conventions.", ACCENT_BLUE),
        ("Accounting Invariance", "Zero-tolerance accounting reconciliation check on every historical record: |Net Flow - (Mobilised - Repurchase)| <= 0.01 Cr.", ACCENT_GREEN),
        ("Exogenous Macro Drivers", "RBI Repo Rate corridor (4.00% pandemic low to 6.50% peak [Feb 2023–Jan 2025], eased to 5.25% at Sep 2026) and NSE Nifty 50 returns.", ACCENT_AMBER),
    ]

    for idx, (p_title, p_desc, col_acc) in enumerate(pillars):
        left = 0.8 + idx * (col_w + gap)
        add_card(slide, left, 1.50, col_w, 2.75)
        tb = slide.shapes.add_textbox(Inches(left + 0.18), Inches(1.68), Inches(col_w - 0.36), Inches(2.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = p_title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col_acc
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_SECONDARY

    # Bottom Section: Relational Architecture & Temporal Safeguard
    add_card(slide, 0.8, 4.45, 11.733, 2.20, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb_bot = slide.shapes.add_textbox(Inches(1.05), Inches(4.60), Inches(11.233), Inches(1.90))
    tf_b = tb_bot.text_frame
    tf_b.word_wrap = True

    p = tf_b.paragraphs[0]
    p.text = "Relational SQLite Storage & Strict Temporal Macro Protection"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY
    p.space_after = Pt(8)

    format_bullet_point(tf_b, "ACID Relational Storage (data/finhype.db)", "Normalized SQLite schema with composite unique constraints on (month, category). Idempotent ingestion ensures corrupt or partial batches trigger complete transaction rollback.", space_after=8)
    format_bullet_point(tf_b, "Zero Synthetic Imputation", "Every single observation represents real AMFI disclosures; no synthetic smoothing, forward filling, or unrecorded discontinuities exist.", space_after=8)
    format_bullet_point(tf_b, "Constant Forward Macro Assumption", "Multi-step forward forecasts strictly enforce Delta Repo = 0 and Nifty Return = 0 for future horizons (h=1..3), preventing speculative central bank policy bets from leaking into forward projections.", space_after=4, body_color=ACCENT_CYAN)

    add_slide_footer(slide, 3, 9, source_note="AMFI Statements, RBI MPC Policy Documents (4.00%–6.50%), NSE Nifty 50; app/data/")


# ── Slide 4: System Architecture (Explicit Flow Arrows & Pipeline Layout) ────
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "End-to-End System Architecture with Strict Temporal Isolation", "System Architecture")

    # 6 Explicit Pipeline Stage Cards with Directional Markers
    stages = [
        ("STAGE 1", "AMFI Ingestion", "Raw Excel/HTML\nReconciliation check\nSQLite finhype.db", BORDER_CYAN, ACCENT_CYAN),
        ("STAGE 2", "Feature Pipeline", "53 engineered indicators\nShifted lags (t-1, t-12)\nRolling volatilities", BORDER_BLUE, ACCENT_BLUE),
        ("STAGE 3", "Forecasting", "4 diverse model families\n24 rolling origins\n576 evaluation rows", BORDER_BLUE, ACCENT_BLUE),
        ("STAGE 4", "Error Calibration", "Genuine OOS errors\nTemporal filter (tau+h<=T)\nq10/q90 & q025/q975", BORDER_AMBER, ACCENT_AMBER),
        ("STAGE 5", "Buffer Sizing", "Max(Base Floor, Deficit)\nComonotonic stress rule\nSingle vs. Stitched", BORDER_GREEN, ACCENT_GREEN),
        ("STAGE 6", "Streamlit UI", "6 interactive tabs\nActive policy selector\nResponsive sub-tabs", BORDER_CYAN, ACCENT_CYAN),
    ]

    card_w = 1.78
    gap = 0.21
    for idx, (s_num, s_title, s_desc, b_color, col_acc) in enumerate(stages):
        left = 0.8 + idx * (card_w + gap)
        add_card(slide, left, 1.50, card_w, 2.65, border_color=b_color)
        tb = slide.shapes.add_textbox(Inches(left + 0.12), Inches(1.65), Inches(card_w - 0.24), Inches(2.35))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = s_num
        p.font.name = FONT_HEADING
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = col_acc

        p_t = tf.add_paragraph()
        p_t.text = s_title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_PRIMARY
        p_t.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = s_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_SECONDARY

        # Add connector arrow between cards (except after last card)
        if idx < len(stages) - 1:
            arrow_left = left + card_w + 0.04
            arr_box = slide.shapes.add_textbox(Inches(arrow_left), Inches(2.65), Inches(0.18), Inches(0.35))
            tf_arr = arr_box.text_frame
            tf_arr.margin_left = tf_arr.margin_right = tf_arr.margin_top = tf_arr.margin_bottom = 0
            p_arr = tf_arr.paragraphs[0]
            p_arr.text = "➔"
            p_arr.font.name = FONT_HEADING
            p_arr.font.size = Pt(14)
            p_arr.font.bold = True
            p_arr.font.color.rgb = ACCENT_BLUE

    # Bottom Split Panel: Mathematical Isolation & Empirical Mutation Proof
    c_w = 5.7
    c_left1 = 0.8
    add_card(slide, c_left1, 4.35, c_w, 2.30, bg_color=CARD_BG, border_color=BORDER_BLUE)
    tb_iso = slide.shapes.add_textbox(Inches(c_left1 + 0.20), Inches(4.50), Inches(c_w - 0.40), Inches(2.0))
    tf_i = tb_iso.text_frame
    tf_i.word_wrap = True
    p = tf_i.paragraphs[0]
    p.text = "Mathematical Temporal Isolation"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(6)
    format_bullet_point(tf_i, "Shifted Operators", "All autoregressive lags (_lag_1, _lag_12) and rolling statistics (3M/6M volatilities) operate exclusively on shifted vectors (.shift(1)). Actual observation Y_t never contaminates features at time t.", space_after=6)
    format_bullet_point(tf_i, "Frozen Transformers", "In Ridge Direct regression, StandardScaler is fit exclusively on training slices (t <= T). Scalers are frozen before test-origin prediction.", space_after=4)

    c_left2 = 6.833
    add_card(slide, c_left2, 4.35, c_w, 2.30, bg_color=CARD_BG, border_color=BORDER_GREEN)
    tb_mut = slide.shapes.add_textbox(Inches(c_left2 + 0.20), Inches(4.50), Inches(c_w - 0.40), Inches(2.0))
    tf_m = tb_mut.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "Automated Future-Mutation Invariance Test"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(6)
    format_bullet_point(tf_m, "End-to-End Production Test", "Production test (tests/test_audit_phase2.py) injects extreme shocks into data after origin T and executes the complete feature, modeling, and interval pipeline.", space_after=6)
    format_bullet_point(tf_m, "Numerical Parity Verified", "Historical point forecasts and prediction intervals match within numerical solver tolerance (+/- 1e-5 Cr), verified across 154 passing automated tests.", space_after=4)

    add_slide_footer(slide, 4, 9, source_note="ARCHITECTURE.md; tests/test_audit_phase2.py; 154 passing unit & integration tests")


# ── Slide 5: Forecasting Models and Evaluation ───────────────────────────────
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Forecasting Suite & Genuine Out-of-Sample Error Calibration", "Forecasting Methodology")

    col_w = 2.75
    gap = 0.244
    models = [
        ("Seasonal-Naive", "Cyclical Benchmark", "Annual 12-month cyclical persistence: Y_hat(T+h) = Y(T+h-12). Captures recurring institutional advance-tax cycles without parameter estimation risk.", ACCENT_CYAN),
        ("ETS State-Space", "Adaptive Smoothing", "Exponential Smoothing with additive errors, damped trend, and seasonal components. Adapts dynamically to gradual structural shifts in fund asset base.", ACCENT_BLUE),
        ("SARIMAX", "Exogenous Regressors", "(p,d,q) x (P,D,Q)_12 Seasonal ARIMA incorporating forward deterministic calendar regressors (advance-tax months, quarter-ends, days in month).", ACCENT_GREEN),
        ("Ridge Direct", "Regularized ML", "L2-regularized linear models with separate direct heads for h in {1, 2, 3}. Internal feature standardization strictly fit on historical training slices.", ACCENT_AMBER),
    ]

    for idx, (m_name, m_tag, m_desc, col_acc) in enumerate(models):
        left = 0.8 + idx * (col_w + gap)
        add_card(slide, left, 1.50, col_w, 2.75)
        tb = slide.shapes.add_textbox(Inches(left + 0.18), Inches(1.68), Inches(col_w - 0.36), Inches(2.45))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = m_name
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col_acc
        
        p_sub = tf.add_paragraph()
        p_sub.text = m_tag.upper()
        p_sub.font.name = FONT_HEADING
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_SECONDARY

    # Bottom Protocol & Calibration Card
    add_card(slide, 0.8, 4.45, 11.733, 2.20, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb_proto = slide.shapes.add_textbox(Inches(1.05), Inches(4.60), Inches(11.233), Inches(1.90))
    tf_p = tb_proto.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    p.text = "24-Origin Expanding Rolling Backtest & Out-of-Sample Calibration Engine"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY
    p.space_after = Pt(8)

    format_bullet_point(tf_p, "Rolling Backtest Window (July 2024 – June 2026)", "24 consecutive monthly origins generating 576 evaluation records (24 origins x 2 categories x 4 models x 3 horizons). Re-fits all models dynamically at each origin T.", space_after=8)
    format_bullet_point(tf_p, "In-Sample Residuals Strictly Prohibited", "Prediction intervals do not use artificial training residuals. Residual error e_(tau,h) = Y_(tau+h) - Y_hat_(tau+h) is admitted to calibration pool if and only if tau + h <= T.", space_after=8)
    format_bullet_point(tf_p, "Calibrated Interval Bounds", "Two-sided 80% (q10 to q90) and 95% (q025 to q975) intervals. Strict assertions enforce monotonicity (l95 <= l80 <= u80 <= u95) and verify point forecast inclusion.", space_after=4, body_color=ACCENT_CYAN)

    add_slide_footer(slide, 5, 9, source_note="app/forecasting/backtest.py, app/forecasting/intervals.py; data/processed/backtest_results.csv")


# ── Slide 6: Backtest Results (Enhanced Hierarchy & Clear Model Identifiers) ─
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Audited Backtest Benchmarks & 3-Month Forward Forecasts", "Empirical Results · Selected Benchmark Summary")

    # Table with Section Header Rows for Visual Hierarchy
    rows_data = [
        # Liquid Fund Section
        ["LIQUID FUNDS (Core Working Capital, AUM: ₹6,25,847 Cr)", "", "", "", "", "", "", ""],
        ["SARIMAX", "1M (h=1)", "₹27,618.17 Cr", "₹35,765.32 Cr", "44.12%", "0.7144", "+14.55%", "★ Winner (Top 1M–3M)"],
        ["Seasonal-Naive", "1M (h=1)", "₹32,321.53 Cr", "₹40,318.61 Cr", "51.63%", "0.8361", "0.00%", "Baseline Benchmark"],
        ["ETS", "1M (h=1)", "₹41,157.21 Cr", "₹50,464.32 Cr", "65.74%", "1.0646", "-27.34%", "Underperforming"],
        ["Ridge Direct", "1M (h=1)", "₹60,698.15 Cr", "₹78,866.64 Cr", "96.96%", "1.5701", "-87.79%", "Underperforming"],
        # Overnight Fund Section
        ["OVERNIGHT FUNDS (Hyper-Velocity Cash, AUM: ₹84,033 Cr)", "", "", "", "", "", "", ""],
        ["SARIMAX", "1M (h=1)", "₹13,940.84 Cr", "₹17,784.25 Cr", "73.62%", "0.9799", "+7.62%", "★ Winner at 1M"],
        ["Seasonal-Naive", "1M (h=1)", "₹15,090.39 Cr", "₹18,116.13 Cr", "79.69%", "1.0607", "0.00%", "Baseline Benchmark"],
        ["Seasonal-Naive", "2M (h=2)", "₹16,076.12 Cr", "₹19,366.68 Cr", "82.08%", "1.1300", "Top 2M", "★ Winner at 2M"],
        ["Seasonal-Naive", "3M (h=3)", "₹15,804.12 Cr", "₹19,157.35 Cr", "82.82%", "1.1109", "Top 3M", "★ Winner at 3M"],
    ]

    col_widths = [Inches(1.6), Inches(1.0), Inches(1.5), Inches(1.5), Inches(1.0), Inches(1.0), Inches(1.4), Inches(2.733)]
    table_shape = slide.shapes.add_table(11, 8, Inches(0.8), Inches(1.42), Inches(11.733), Inches(2.45))
    table = table_shape.table

    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    # Explicit row heights
    table.rows[0].height = Inches(0.25)
    for r in range(1, 11):
        table.rows[r].height = Inches(0.22)

    headers = ["Model", "Horizon", "MAE (₹ Cr)", "RMSE (₹ Cr)", "WAPE", "MASE", "Skill Score", "Benchmark Status"]
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG_LIGHT
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.margin_left = Inches(0.04)
        cell.margin_right = Inches(0.04)
        cell.text = h_text
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_HEADING
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.alignment = PP_ALIGN.CENTER

    # Merge Row 1 across columns 0..7 for Liquid Funds Section Header
    table.cell(1, 0).merge(table.cell(1, 7))
    c1 = table.cell(1, 0)
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_HIGHLIGHT
    c1.margin_top = Inches(0.02)
    c1.margin_bottom = Inches(0.02)
    c1.margin_left = Inches(0.08)
    c1.margin_right = Inches(0.04)
    c1.text = rows_data[0][0]
    p1 = c1.text_frame.paragraphs[0]
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(9.5)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN
    p1.alignment = PP_ALIGN.LEFT

    # Merge Row 6 across columns 0..7 for Overnight Funds Section Header
    table.cell(6, 0).merge(table.cell(6, 7))
    c6 = table.cell(6, 0)
    c6.fill.solid()
    c6.fill.fore_color.rgb = CARD_HIGHLIGHT
    c6.margin_top = Inches(0.02)
    c6.margin_bottom = Inches(0.02)
    c6.margin_left = Inches(0.08)
    c6.margin_right = Inches(0.04)
    c6.text = rows_data[5][0]
    p6 = c6.text_frame.paragraphs[0]
    p6.font.name = FONT_HEADING
    p6.font.size = Pt(9.5)
    p6.font.bold = True
    p6.font.color.rgb = ACCENT_AMBER
    p6.alignment = PP_ALIGN.LEFT

    # Populate data rows
    for row_idx, r_data in enumerate(rows_data):
        if row_idx in (0, 5):
            continue
        table_r = row_idx + 1
        for col_idx, val in enumerate(r_data):
            cell = table.cell(table_r, col_idx)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else SURFACE_COLOR
            cell.text = str(val)
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_BODY
            p.font.size = Pt(9)
            if "★ Winner" in str(val) or "Top" in str(val):
                p.font.bold = True
                p.font.color.rgb = ACCENT_GREEN
            elif "-" in str(val) and "%" in str(val):
                p.font.color.rgb = ACCENT_RED
            else:
                p.font.color.rgb = TEXT_PRIMARY if col_idx == 0 else TEXT_SECONDARY
            p.alignment = PP_ALIGN.CENTER

    # Bottom Cards: Key Empirical Takeaway and Forward Net Flow Projections
    c_top = 4.25
    c_h = 2.40
    c_left1 = 0.8
    c_w1 = 5.7
    add_card(slide, c_left1, c_top, c_w1, c_h)
    tb_ins = slide.shapes.add_textbox(Inches(c_left1 + 0.18), Inches(c_top + 0.12), Inches(c_w1 - 0.36), Inches(c_h - 0.24))
    tf_i = tb_ins.text_frame
    tf_i.word_wrap = True
    tf_i.margin_left = tf_i.margin_right = tf_i.margin_top = tf_i.margin_bottom = 0
    p = tf_i.paragraphs[0]
    p.text = "Empirical Model Hierarchy Takeaway"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(4)
    format_bullet_point(tf_i, "Liquid Funds (SARIMAX Dominant)", "SARIMAX outperforms across all horizons (1M: ₹27,618 Cr [+14.55% Skill], 2M: ₹31,620 Cr, 3M: ₹33,087 Cr) due to deterministic advance-tax exogenous regressors.", space_after=4, font_size_title=10, font_size_body=9.5)
    format_bullet_point(tf_i, "Overnight Funds (Horizon Split)", "SARIMAX wins at Horizon 1 (MAE ₹13,941 Cr, +7.62% Skill). Beyond 1 month, Seasonal-Naive wins (MAE ₹16,076 Cr at 2M, ₹15,804 Cr at 3M) as macro relationships break down under 545% turnover.", space_after=2, font_size_title=10, font_size_body=9.5)

    c_left2 = 6.833
    c_w2 = 5.7
    add_card(slide, c_left2, c_top, c_w2, c_h)
    tb_fwd = slide.shapes.add_textbox(Inches(c_left2 + 0.18), Inches(c_top + 0.12), Inches(c_w2 - 0.36), Inches(c_h - 0.24))
    tf_f = tb_fwd.text_frame
    tf_f.word_wrap = True
    tf_f.margin_left = tf_f.margin_right = tf_f.margin_top = tf_f.margin_bottom = 0
    p = tf_f.paragraphs[0]
    p.text = "Forward Projections by Selected Model (Origin: Sep 2026)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(4)
    format_bullet_point(tf_f, "Liquid Funds (SARIMAX for h=1..3)", "Oct 2026: +₹75,820.52 Cr (rebound) -> Nov 2026: -₹8,979.98 Cr -> Dec 2026: -₹50,618.62 Cr (corporate tax drainage, lower 80% bound: -₹89,235.07 Cr).", space_after=4, font_size_title=10, font_size_body=9.5)
    format_bullet_point(tf_f, "Overnight Funds (SARIMAX 1M / S-Naive 2M,3M)", "Oct 2026: +₹16,981.99 Cr (SARIMAX) -> Nov 2026: -₹37,624.53 Cr (S-Naive persistence of Nov 2025 outlier) -> Dec 2026: +₹254.25 Cr (S-Naive).", space_after=2, font_size_title=10, font_size_body=9.5)

    add_slide_footer(slide, 6, 9, source_note="data/processed/backtest_metrics_summary.csv (Full 32-series evaluation); forward_forecasts_3m.csv")


# ── Slide 7: Dynamic Cash-Buffer Engine (High-Impact Sizing Callouts) ────────
def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Dynamic Cash Buffer Engine & Dual Overnight Policies", "Cash Buffer Engine")

    # High-Impact Formulation Banner
    add_card(slide, 0.8, 1.45, 11.733, 1.10, bg_color=CARD_BG, border_color=BORDER_BLUE)
    tb_form = slide.shapes.add_textbox(Inches(1.05), Inches(1.52), Inches(11.233), Inches(0.95))
    tf_fo = tb_form.text_frame
    tf_fo.word_wrap = True
    p = tf_fo.paragraphs[0]
    p.text = "Sizing Formula:  Recommended Buffer = max(Base Operational Floor,  Peak Cumulative Downside Deficit)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(3)
    p2 = tf_fo.add_paragraph()
    p2.text = (
        "Non-Monotonic Preservation: Peak Deficit = max_{k=1..h} [max(0, -Sum_{j=1..k} L_hat_{T+j})]. "
        "Early liquidity shortfalls cannot be masked or erased by later projected cash inflows."
    )
    p2.font.name = FONT_BODY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SECONDARY

    # Left Card: Liquid Fund Recommendations
    add_card(slide, 0.8, 2.70, 5.7, 3.95)
    tb_l = slide.shapes.add_textbox(Inches(1.05), Inches(2.85), Inches(5.2), Inches(3.65))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "Liquid Fund Buffer (AUM: ₹6,25,847 Cr | SARIMAX)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)

    format_bullet_point(tf_l, "1-Month (Oct 2026)", "Buffer: ₹31,292.35 Cr (5.00% AUM). Point forecast is +₹75,821 Cr, downside deficit is ₹0; governed entirely by base operational floor.", space_after=8)
    format_bullet_point(tf_l, "2-Month (Oct–Nov 2026)", "Buffer: ₹31,292.35 Cr (5.00% AUM). Cumulative downside remains positive (+₹14,162 Cr); governed by operational floor.", space_after=8)
    format_bullet_point(tf_l, "3-Month Moderate (80%)", "Buffer: ₹75,072.74 Cr (12.00% AUM). Cumulative downside deficit expands to ₹75,073 Cr ahead of December corporate advance-tax drainage.", space_after=8, body_color=ACCENT_AMBER)
    format_bullet_point(tf_l, "3-Month Conservative (95%)", "Buffer: ₹106,161.75 Cr (16.96% AUM). Sized to absorb severe 2.5th percentile downside tax drain (-₹94,506 Cr bound).", space_after=4, body_color=ACCENT_RED)

    # Right Card: Overnight Fund Recommendations & Outlier Disclosures
    add_card(slide, 6.833, 2.70, 5.7, 3.95)
    tb_r = slide.shapes.add_textbox(Inches(7.083), Inches(2.85), Inches(5.2), Inches(3.65))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Overnight Fund Buffer (AUM: ₹84,033 Cr | Dual Policies)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)

    format_bullet_point(tf_r, "1-Month Standalone (SARIMAX)", "Moderate (80%): ₹10,366.04 Cr (12.34% AUM) | Conservative (95%): ₹17,606.86 Cr (20.95% AUM). Sized to absorb 1M turnover.", space_after=8)
    format_bullet_point(tf_r, "3-Month Single Baseline (S-Naive)", "Moderate: ₹74,130.35 Cr (88.22% AUM) | Conservative: ₹118,049.89 Cr (140.48% AUM). Replicates Nov 2025 outlier (-₹37,625 Cr).", space_after=8)
    format_bullet_point(tf_r, "3-Month Stitched (SARIMAX+S-Naive)", "Moderate: ₹90,640.30 Cr (107.86% AUM) | Conservative: ₹129,620.28 Cr (154.25% AUM). Matches empirical champion at each step.", space_after=8)
    format_bullet_point(tf_r, "Why Buffer Exceeds 100% of AUM", "Overnight monthly gross turnover exceeds 540% of AUM. Comonotonic compounding of Nov 2025 outlier models extreme gross redemption distress, preserved transparently without artificial caps.", space_after=4, body_color=ACCENT_CYAN)

    add_slide_footer(slide, 7, 9, source_note="app/services/cash_buffer.py, app/services/forward_forecasting.py; Origin: September 2026")


# ── Slide 8: Live Dashboard Demonstration (High-Fidelity UI Schematic) ───────
def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Interactive Streamlit Platform: Decision Support Architecture", "Dashboard Demonstration")

    # Top Tab Navigation Bar Schematic
    add_card(slide, 0.8, 1.45, 11.733, 0.65, bg_color=CARD_BG_LIGHT, border_color=BORDER_BLUE)
    tb_nav = slide.shapes.add_textbox(Inches(0.95), Inches(1.52), Inches(11.433), Inches(0.50))
    tf_nav = tb_nav.text_frame
    p_nav = tf_nav.paragraphs[0]
    p_nav.text = "TABS:  1. Exec Summary   |   2. Historical & AUM   |   3. Forward Forecasts [Active]   |   4. Buffer Sizing   |   5. Model Benchmark   |   6. Governance"
    p_nav.font.name = FONT_HEADING
    p_nav.font.size = Pt(11)
    p_nav.font.bold = True
    p_nav.font.color.rgb = ACCENT_CYAN

    # Left Panel: Tab 3 Forward Forecast Features
    col_w = 5.7
    c1_left = 0.8
    add_card(slide, c1_left, 2.25, col_w, 4.25)
    tb1 = slide.shapes.add_textbox(Inches(c1_left + 0.22), Inches(2.40), Inches(col_w - 0.44), Inches(3.95))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Tab 3: Forward Forecasts Architecture"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.space_after = Pt(10)

    format_bullet_point(tf1, "Active Policy Trajectory Selector", "Default dropdown option directly reflects the exact model sequence deployed in Tab 4's cash buffer engine, ensuring complete cross-tab parity.", space_after=10)
    format_bullet_point(tf1, "Standalone Model Exploration", "Selecting candidate models (SARIMAX, S-Naive, ETS, Ridge) renders pure unconstrained trajectories with contextual notice banners.", space_after=10)
    format_bullet_point(tf1, "Responsive 3-View Sub-Tabs", "Structured sub-tabs (Complete Summary, Moderate 80%, Conservative 95%) prevent horizontal clipping on 1280px laptop screens.", space_after=6)

    # Right Panel: Tab 4 Treasury Sizing & Progression Controls
    c2_left = 6.833
    add_card(slide, c2_left, 2.25, col_w, 4.25)
    tb2 = slide.shapes.add_textbox(Inches(c2_left + 0.22), Inches(2.40), Inches(col_w - 0.44), Inches(3.95))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Tab 4: Treasury Buffer Controls"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(10)

    format_bullet_point(tf2, "Real-Time Treasury Controls", "Interactive sliders for Base Operational Floor %, Planning Horizon (1M-3M), and Risk Tolerance (Moderate 80% vs Conservative 95%).", space_after=10)
    format_bullet_point(tf2, "Deficit Progression Audit Table", "Step-by-step table explicitly details Target Month, Model, Point Forecast, Downside Bound, and Cumulative Deficit.", space_after=10)
    format_bullet_point(tf2, "Policy Comparison Expander", "Live side-by-side comparison of Overnight Single-Model Baseline vs. Step-Wise Stitched Alternative.", space_after=6)

    # Bottom Launch Command Banner
    add_card(slide, 0.8, 6.00, 11.733, 0.65, bg_color=CARD_BG, border_color=BORDER_COLOR)
    tb_cmd = slide.shapes.add_textbox(Inches(1.05), Inches(6.08), Inches(11.233), Inches(0.48))
    tf_c = tb_cmd.text_frame
    p = tf_c.paragraphs[0]
    p.text = "One-Command Launch:  .\\run.ps1 -OnlyDashboard  (PowerShell)  |  .\\run.bat --only-dashboard  (CMD)  |  Instant Cache Startup < 1.0s"
    p.font.name = FONT_HEADING
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    add_slide_footer(slide, 8, 9, source_note="app/dashboard.py; DEMO_GUIDE.md; Streamlit Web Platform (UI Schematic)")


# ── Slide 9: Limitations, Governance, and Conclusion (Strict Separation) ─────
def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Methodological Boundaries, Governance & Production Conclusion", "Governance & Conclusion")

    # Left Card: Transparent Disclosures (What We Do NOT Claim)
    add_card(slide, 0.8, 1.50, 5.7, 5.15)
    tb_l = slide.shapes.add_textbox(Inches(1.05), Inches(1.68), Inches(5.2), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "Methodological Boundaries & Disclosures"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p.space_after = Pt(12)

    format_bullet_point(tf_l, "Comonotonic Stress Scenario", "Summing monthly downside bounds assumes perfect positive tail dependence (consecutive worst-case shocks). It is an advisory stress scenario, not a calibrated joint multi-period probability interval.", space_after=11)
    format_bullet_point(tf_l, "Sample Size Discreteness (N=24)", "Evaluating 24 backtest origins limits statistical resolution; each observation is ~4.17% of the empirical distribution. Nominal 80%/95% levels are empirical calibration heuristics, not mathematical proofs.", space_after=11)
    format_bullet_point(tf_l, "Monthly Aggregation vs. Intraday Runs", "AMFI data is monthly aggregate. Intraday redemption spikes occurring on the 15th of tax months must be buffered via operational floors, sponsor lines, and repo credit facilities.", space_after=11)
    format_bullet_point(tf_l, "Overnight Buffer >100% Reality", "Reflects acute comonotonic stress under 545% monthly turnover, not a claim that the fund literally holds that amount in physical cash.", space_after=4, body_color=ACCENT_CYAN)

    # Right Card: Production Roadmap & Final Verdict
    add_card(slide, 6.833, 1.50, 5.7, 5.15)
    tb_r = slide.shapes.add_textbox(Inches(7.083), Inches(1.68), Inches(5.2), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Production Roadmap & Release Verdict"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    format_bullet_point(tf_r, "Roadmap 1: Vine Copula Dependencies", "Replace comonotonic summation with calibrated parametric or vine copulas to model non-linear inter-temporal dependence across multi-month planning horizons.", space_after=11)
    format_bullet_point(tf_r, "Roadmap 2: Daily Transaction Integration", "Connect ingestion pipeline to AMC transfer agency feeds (CAMS/KFintech) for intraday RTGS flow monitoring and tax calendar weighting.", space_after=11)
    format_bullet_point(tf_r, "Roadmap 3: Dynamic Macro Scenarios", "Incorporate live RBI MPC probability distributions and liquidity adjustment facility (LAF) balances for interactive rate-shock stress testing.", space_after=11)
    format_bullet_point(tf_r, "System Readiness Sign-Off", "FINHYPE is complete, mathematically reconciled, and verified across 154 passing unit and integration tests. Certified ready for Zeckathon 2026 defense.", space_after=4, body_color=ACCENT_CYAN)

    add_slide_footer(slide, 9, 9, source_note="LIMITATIONS.md, ARCHITECTURE.md, RELEASE_CHECKLIST.md; Final Certified Release")


def main():
    output_path = Path("C:/Users/Alok/Desktop/FINHYPE/FINHYPE_Zeckathon_2026.pptx")
    print(f"Generating polished 9-slide presentation (v2): {output_path}...")
    prs = create_base_presentation()

    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)

    prs.save(str(output_path))
    print(f"Successfully generated {output_path} ({len(prs.slides)} slides)!")


if __name__ == "__main__":
    main()
